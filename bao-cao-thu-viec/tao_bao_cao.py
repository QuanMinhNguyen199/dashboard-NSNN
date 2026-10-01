from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
from datetime import datetime

OUT = Path(__file__).resolve().parent
SOURCE = next(p for p in Path('C:/Users/admin/Downloads').glob('*.docx') if 'THEHEGEO_HDTV' in p.name and '(2) (1)' in p.name)
weeks = [
('17/08/2026', '21/08/2026', [
('Tìm hiểu quy định công ty THEHEGEO, Outline, Plane và cách làm việc của vị trí BQA.',
 'Đã đọc quy định làm việc; truy cập thành công Outline và Plane; tìm hiểu cách tra cứu tài liệu, ghi nhận yêu cầu và phối hợp công việc. Minh chứng: báo cáo onboarding tuần 17–21/08.'),
('Hoàn thành các nội dung học nền tảng về Business Analysis; tìm hiểu quy trình từ yêu cầu đến kiểm thử.',
 'Hoàn thành 3 nội dung học, tổng 9 giờ: IT Business Analysts (2 giờ), Business Analysis Certification Course (3 giờ), LEARN CBAP (4 giờ). Nắm vai trò BA, User Story, Acceptance Criteria, quản lý yêu cầu và Verify/Validate.'),
('Đọc yêu cầu HKD2026, DN360 và thực hành kiểm thử Agent Thuế trên GeoAgent UAT.',
 'Đã đọc 2 bộ yêu cầu, xây dựng scenario kiểm thử Agent Thuế, viết 20 test case, thực hiện 10 test và lập 5 bug report. Ghi nhận các vấn đề cần trao đổi với mentor về nguồn dữ liệu, trích dẫn, cách tính và ngưỡng rủi ro.'),
]),
('24/08/2026', '28/08/2026', [
('Tìm hiểu Calendar/LibreBooking; phân tích chức năng đặt phòng, quản lý tài nguyên và phạm vi phân quyền.',
 'Lập tài liệu chức năng và đánh giá khả năng đáp ứng, phân biệt chức năng có thể kế thừa với phần THG cần bổ sung. Minh chứng: tài liệu LibreBooking ngày 24/08.'),
('Xây dựng đặc tả, kế hoạch triển khai và tài liệu bàn giao cho các chức năng đặt phòng, tổng quan, danh mục phòng và báo cáo.',
 'Có bộ tài liệu spec, implement-plan và handoff theo chức năng; mô tả luồng, API sử dụng, trường hợp lỗi và tiêu chí kiểm tra. Minh chứng: các thư mục docs của Calendar.'),
('Tham gia dựng và rà soát giao diện Calendar theo prototype; kiểm tra luồng tìm phòng, đặt lịch, sửa/hủy và chọn người tham dự.',
 'Hình thành các màn nghiệp vụ và bộ thành phần giao diện dùng chung; bổ sung quản lý phòng, thiết bị và báo cáo dành cho quản trị. Minh chứng: cập nhật Calendar từ 26–28/08.'),
]),
('31/08/2026', '04/09/2026', [
('Rà soát và hoàn thiện thao tác đặt phòng trực tiếp trên lịch, cách thể hiện lịch nhiều ngày và lịch toàn công ty.',
 'Cập nhật luồng đặt phòng trên lưới lịch và cách hiển thị sự kiện qua nhiều ngày. Minh chứng: đặc tả calendar-booking và cập nhật ngày 03/09.'),
('Hiệu chỉnh luồng phê duyệt và biểu mẫu: bổ sung lý do từ chối, rà soát thao tác xóa và đường dẫn của từng form.',
 'Form có URL riêng; luồng từ chối ghi nhận lý do và thao tác xóa được kiểm soát rõ hơn. Minh chứng: cập nhật Calendar ngày 04/09.'),
]),
('07/09/2026', '11/09/2026', [
('Tiếp tục rà soát Calendar: biểu mẫu đặt phòng, quản lý tài nguyên, trạng thái chờ duyệt và nội dung thông báo.',
 'Thống nhất thành phần dùng chung; hiệu chỉnh quy tắc thời lượng, điều kiện check-in và trạng thái phê duyệt. Minh chứng: cập nhật Calendar từ 07–09/09.'),
('Kiểm tra cách tổng hợp báo cáo và hiển thị lịch; chuẩn hóa nhãn, tiêu đề bảng và kỳ thống kê.',
 'Điều chỉnh kỳ báo cáo và số lượng lịch chờ duyệt theo tháng; đồng bộ cách diễn đạt trên giao diện. Minh chứng: tài liệu reports và các cập nhật ngày 08–09/09.'),
('Xây dựng prototype Dashboard Thu NSNN Hà Nội và tài liệu phục vụ trình diễn, góp ý nghiệp vụ.',
 'Dựng 4 khu vực: Tổng quan, Phân tích thu, Chi tiết phường/xã, So sánh nâng cao; bổ sung bản xem thử khung nhúng và bố cục màn hình hẹp. Số liệu ở giai đoạn này là mô phỏng.'),
]),
('14/09/2026', '18/09/2026', [
('Hoàn thiện tài liệu nghiệp vụ và thiết kế Dashboard NSNN; rà soát biểu đồ cơ cấu thu, bản đồ, bộ lọc và luồng xem chi tiết.',
 'Có báo cáo prototype ngày 14/09, tài liệu BA và thiết kế dashboard; cập nhật bố cục, typography và bộ lọc mobile. Phạm vi địa bàn gồm 126 phường/xã.'),
('Phân tích tài liệu và dữ liệu bàn giao; đối chiếu nhu cầu báo cáo theo cơ quan thuế, ngành nghề, dự toán và nguồn TMS/TTR.',
 'Bổ sung khung phân tích TMS và bộ lọc cấp quản lý; lập roadmap, chỉ rõ nguồn còn thiếu và nội dung chưa đủ điều kiện xác nhận. Minh chứng: cập nhật 16–18/09 và ROADMAP.md.'),
('Đưa dữ liệu tổng hợp đã nhận vào prototype; phân biệt dữ liệu nguồn với các chỉ tiêu còn mô phỏng.',
 'Tổng hợp danh bạ 473.618 bản ghi và dự toán của 126 phường/xã theo tài liệu ngày 18/09. Phần số thu chi tiết, kết quả kiểm tra và chỉ tiêu sai số dự báo vẫn có giới hạn cần xác nhận.'),
]),
('21/09/2026', '25/09/2026', [
('Tiếp tục hoàn thiện Dashboard theo dữ liệu TMS/TTR và nhu cầu phân tích thu; rà soát trạng thái tải, bảng phân rã và cách đọc chỉ tiêu.',
 'Cập nhật khu vực TMS/TTR, bảng cơ quan thuế quản lý địa bàn và phân tích thu. Minh chứng: các cập nhật NSNN ngày 21–23/09.'),
('Rà soát giao diện trên desktop/mobile; điều chỉnh bộ lọc, KPI, biểu đồ xu hướng và danh sách người nộp thuế.',
 'Cải thiện bộ lọc mobile, thao tác cuộn và cách trình bày KPI; bổ sung công cụ/ảnh kiểm tra giao diện. Minh chứng: các cập nhật NSNN ngày 24–25/09.'),
('Rà soát cách phân loại ngành nghề và phạm vi hiển thị của danh sách người nộp thuế trên dashboard.',
 'Điều chỉnh danh sách xếp hạng theo ngành, xử lý nhóm chưa phân loại trong phạm vi hiển thị. Minh chứng: cập nhật chức năng danh sách người nộp thuế ngày 25/09.'),
]),
('28/09/2026', '01/10/2026', [
('Hiệu chỉnh nhận diện Calendar; tiếp tục chuẩn hóa nội dung và cách thể hiện số liệu trên Dashboard NSNN.',
 'Cập nhật màu nhận diện, thẻ phòng và bảng của Calendar; thống nhất nhãn, tiêu đề và định dạng tiền trên NSNN. Minh chứng: các cập nhật từ 28–30/09.'),
('Phân tích tài liệu khảo sát và xây dựng prototype tax-ops cho tác nghiệp nội bộ; thể hiện luồng báo cáo theo vai trò.',
 'Có các nhóm chức năng công việc, nợ, rủi ro, hoàn thuế, báo cáo và dữ liệu; phân biệt màn cán bộ với màn lãnh đạo khi gửi/duyệt báo cáo. Đây là prototype, chưa xác nhận quy trình vận hành production.'),
('Chuẩn hóa design system từ prototype tax-ops, chuẩn bị tài liệu và bộ thành phần để đưa vào Figma.',
 'Chuẩn bị 38 màu, 11 text styles, 20 icon và 16 nhóm component; preview đã kiểm tra ở 1440px và 390px. Bộ import đã tạo; chưa ghi vào Figma do hết lượt kết nối MCP, còn chờ import và kiểm tra canvas.'),
]),
]

d = Document(SOURCE)
original = d.paragraphs
def replace(p, value):
    if p.runs:
        p.runs[0].text = value
        for r in p.runs[1:]: r.text = ''
    else: p.add_run(value)

replace(original[2], 'Họ tên: Nguyễn Minh Quân                         Ngày sinh: ......................................')
replace(original[3], 'Chức danh: Nhân viên BQA                       Đơn vị: ...........................................')
replace(original[4], 'Thời gian báo cáo: Từ ngày 17/08/2026 đến ngày 01/10/2026')
replace(original[5], 'Người hướng dẫn: Lê Quang Anh                  Chức danh: .....................................')
# Rebuild only section A; keep the employer's confirmation and assessment sections.
start = original[8]._p
end = original[20]._p
node = start.getnext()
while node is not end:
    nxt = node.getnext()
    node.getparent().remove(node)
    node = nxt

def add_p(value='',bold=False,size=11.5,break_before=False):
    p=d.add_paragraph()
    p.paragraph_format.space_after=Pt(5)
    p.paragraph_format.space_before=Pt(5 if bold else 0)
    p.paragraph_format.line_spacing=1.08
    p.paragraph_format.keep_with_next=bold
    p.paragraph_format.page_break_before=break_before
    r=p.add_run(value);r.bold=bold;r.font.size=Pt(size)
    end.addprevious(p._p)
    return p

def cell_text(cell,value,bold=False):
    cell.text=''
    p=cell.paragraphs[0]
    p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.space_before=Pt(4)
    p.paragraph_format.line_spacing=1.05
    r=p.add_run(value);r.bold=bold;r.font.size=Pt(11)

add_p('Báo cáo chỉ ghi nhận ngày làm việc từ thứ Hai đến thứ Sáu; tuần 7 ghi nhận đến ngày 01/10/2026. Phần đánh giá, xác nhận cuối kỳ dành cho người hướng dẫn và đơn vị.',size=10.5)
for idx,(begin,finish,rows) in enumerate(weeks,1):
    short=lambda value: datetime.strptime(value, '%d/%m/%Y').strftime('%d/%m').lstrip('0').replace('/0','/')
    add_p(f'Tuần {idx} ({short(begin)}–{short(finish)})',bold=True,break_before=idx in (2,4,6))
    table=d.add_table(rows=1,cols=2)
    table.autofit=False
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'808080');borders.append(border)
    table._tbl.tblPr.append(borders)
    table.columns[0].width=Cm(7.5);table.columns[1].width=Cm(9.5)
    cell_text(table.rows[0].cells[0],'NỘI DUNG CÔNG VIỆC',True)
    cell_text(table.rows[0].cells[1],'KẾT QUẢ THỰC HIỆN / MINH CHỨNG',True)
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for c in table.rows[0].cells:
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EFEFEF');c._tc.get_or_add_tcPr().append(sh)
    for left,right in rows:
        cells=table.add_row().cells
        cell_text(cells[0],left);cell_text(cells[1],right)
    for row in table.rows:
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
    end.addprevious(table._tbl)
    add_p()

add_p('2. Những yêu cầu khác trong thời gian thử việc',bold=True)
for item in [
    'Tiếp tục cập nhật đặc tả và tiêu chí kiểm tra theo góp ý của người hướng dẫn; bảo đảm tài liệu và prototype thống nhất.',
    'Theo dõi lỗi, ghi nhận kết quả xử lý và lưu minh chứng; phân biệt rõ chức năng đã triển khai, phần mô phỏng và hạng mục còn chờ dữ liệu.',
    'Hoàn thiện bàn giao tài liệu, prototype và design system; phối hợp xác nhận nghiệp vụ trước khi sử dụng chính thức.'
]: add_p('– '+item,size=11)

# Assessment remains blank, rather than inheriting the template's printed × as selections.
for p in d.paragraphs:
    if p.text.startswith('×'):replace(p,p.text.replace('×','☐',1))
original[20].paragraph_format.page_break_before=True
original[23].paragraph_format.page_break_before=False
assessment=False
for p in list(d.paragraphs):
    if p._p is end: assessment=True
    if assessment:
        if not p.text.strip() and not p._p.xpath('.//w:drawing'):
            p._p.getparent().remove(p._p)
            continue
        p.paragraph_format.space_before=Pt(3)
        p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.line_spacing=1
for sec in d.sections:
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.left_margin=Cm(2);sec.right_margin=Cm(2);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8)
    sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
    for footer in (sec.footer,sec.first_page_footer,sec.even_page_footer):
        for child in list(footer._element): footer._element.remove(child)
        footer._element.append(OxmlElement('w:p'))
for p in d.paragraphs:
    for r in p.runs:
        r.font.name='Times New Roman'
        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
for table in d.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:r.font.name='Times New Roman'
d.core_properties.author='Nguyễn Minh Quân'
d.core_properties.title='Báo cáo thử việc Nguyễn Minh Quân — 17/08/2026 đến 01/10/2026'
d.core_properties.subject='Tổng hợp Calendar, Dashboard NSNN và tax-ops'
dest=OUT/'Bao-cao-thu-viec_Nguyen-Minh-Quan_17-08_den_01-10-2026.docx'
d.save(dest)
# Sanity checks against the actual saved document.
check=Document(dest)
content='\n'.join(p.text for p in check.paragraphs)+'\n'+'\n'.join(c.text for t in check.tables for row in t.rows for c in row.cells)
assert all(f'Tuần {i} (' in content for i in range(1,8))
assert all(datetime.strptime(date,'%d/%m/%Y').weekday()<5 for begin,finish,_ in weeks for date in (begin,finish))
assert all(not ''.join(f._element.itertext()).strip() for s in check.sections for f in (s.footer,s.first_page_footer,s.even_page_footer))
assert 'Ví dụ:' not in content and 'Gợi ý:' not in content
assert len(check.tables)==9, len(check.tables)
assert all(x in content for x in ['B. XÁC NHẬN','C. ĐÁNH GIÁ THỬ VIỆC','Nguyễn Minh Quân','Quang Anh','Nhân viên BQA'])
(OUT/'noi-dung-bao-cao.json').write_text(json.dumps(weeks,ensure_ascii=False,indent=2),encoding='utf-8')
print(dest)
