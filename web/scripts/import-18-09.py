"""Tạo tập dữ liệu tổng hợp an toàn từ bộ tài liệu ngày 18-09.

Script chỉ xuất số tổng hợp phục vụ prototype. Mã số thuế và tên người nộp
thuế trong danh bạ không bao giờ được ghi vào repository.
"""

from __future__ import annotations

from collections import Counter
import json
from pathlib import Path
import re
import unicodedata

import openpyxl
import xlrd


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "18-09"
OUTPUT = ROOT / "web" / "src" / "data" / "official-18-09.json"
WARDS = ROOT / "web" / "src" / "domain" / "wards.json"


def key(value: object) -> str:
    # Bỏ tiền tố trên chuỗi còn dấu. Nếu bỏ sau khi khử dấu, tên riêng
    # "Phương Liệt" sẽ bị hiểu nhầm chữ "Phương" thành tiền tố "Phường".
    text = re.sub(r"^(phường|xã)\s+", "", str(value or "").strip().lower())
    text = unicodedata.normalize("NFD", text)
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")


def number(value: object) -> int:
    """Đổi triệu đồng trong file nguồn thành đồng nguyên."""
    return round(float(value or 0) * 1_000_000)


def location_plans() -> list[dict[str, object]]:
    wards = json.loads(WARDS.read_text(encoding="utf-8"))
    ward_by_name = {key(row["location_name"]): row for row in wards}
    path = SOURCE / "2. Du toan 126 XP nam 2026.xlsx"
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows: list[dict[str, object]] = []
    component_columns = {
        "nonStateBusiness": 5,
        "registrationFee": 10,
        "environmentTax": 13,
        "agriculturalLandTax": 14,
        "nonAgriculturalLandTax": 15,
        "personalIncomeTax": 16,
        "fees": 20,
        "landRent": 22,
        "landUse": 26,
        "landCompensation": 32,
        "publicLand": 33,
        "otherRevenue": 34,
    }
    for row_no in [*range(10, 61), *range(62, 137)]:
        raw_name = sheet.cell(row_no, 2).value
        ward = ward_by_name.get(key(raw_name))
        if not ward:
            raise ValueError(f"Không nối được địa bàn: {raw_name!r}")
        rows.append(
            {
                "id": ward["location_code"],
                "name": ward["location_name"],
                "plan": number(sheet.cell(row_no, 3).value),
                "planExcludingLandUse": number(sheet.cell(row_no, 4).value),
                "components": {
                    name: number(sheet.cell(row_no, column).value)
                    for name, column in component_columns.items()
                },
            }
        )
    workbook.close()
    if len(rows) != 126:
        raise ValueError(f"Cần đủ 126 địa bàn, nhận được {len(rows)}")
    return rows


def city_history() -> list[dict[str, int]]:
    path = SOURCE / "1. Du toan + So thu 2015-2026 (Toan HN).xlsx"
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    result = [{"year": 2015, "actual": number(sheet.cell(5, 3).value)}]
    column = 4
    for year in range(2016, 2027):
        result.append(
            {
                "year": year,
                "plan": number(sheet.cell(5, column).value),
                "actualOrEstimate": number(sheet.cell(5, column + 1).value)
                if sheet.cell(5, column + 1).value is not None
                else 0,
            }
        )
        column += 2
    workbook.close()
    return result


def forecast_2026() -> dict[str, object]:
    path = SOURCE / "3. Tham khai uoc thuc hien 2026 (tung thang cuoi nam).xls"
    workbook = xlrd.open_workbook(path, on_demand=True)
    sheet = workbook.sheet_by_index(0)
    # Dòng 10 trong Excel tương ứng index 9: tổng thu I+II.
    row = sheet.row_values(9)
    result = {
        "actualFirstSixMonths": number(row[4]),
        "monthlyEstimate": [
            {"month": month, "amount": number(row[column])}
            for month, column in zip(range(7, 13), range(5, 11), strict=True)
        ],
        "annualEstimate": number(row[11]),
    }
    workbook.release_resources()
    return result


def directory_aggregates() -> dict[str, object]:
    path = SOURCE / "5. MAP Danh ba theo NNKD.xlsx"
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["DN Trongdiem"]
    iterator = sheet.iter_rows(values_only=True)
    header = None
    for row in iterator:
        if "Mã số thuế" in row:
            header = [str(value or "").strip() for value in row]
            break
    if not header:
        raise ValueError("Không tìm thấy header danh bạ")
    positions = {name: header.index(name) for name in [
        "CQT",
        "Mã số thuế",
        "ĐVQL NNT",
        "PL_NNKD",
        "Xã/ phường đặt trụ sở chính",
    ]}
    industry: Counter[str] = Counter()
    tax_office_code: Counter[str] = Counter()
    managing_unit: Counter[str] = Counter()
    location: Counter[str] = Counter()
    total = 0
    for row in iterator:
        if not row[positions["Mã số thuế"]]:
            continue
        total += 1
        industry[str(row[positions["PL_NNKD"]] or "Chưa xác định").strip()] += 1
        raw_code = str(row[positions["CQT"]] or "").strip()
        # Excel có thể đọc mã bốn chữ số thành số; trả lại số 0 đầu để nối với
        # danh mục cơ quan thuế của dashboard.
        office_code = raw_code.removesuffix(".0").zfill(4) if raw_code else "Chưa xác định"
        tax_office_code[office_code] += 1
        managing_unit[str(row[positions["ĐVQL NNT"]] or "Chưa xác định").strip()] += 1
        location[str(row[positions["Xã/ phường đặt trụ sở chính"]] or "Chưa xác định").strip()] += 1
    workbook.close()
    return {
        "total": total,
        "byIndustry": dict(industry),
        "byTaxOfficeCode": dict(tax_office_code),
        "byManagingUnit": dict(managing_unit),
        "byLocation": dict(location),
    }


def main() -> None:
    location_rows = location_plans()
    payload = {
        "meta": {
            "sourceSet": "Tài liệu nghiệp vụ 18-09",
            # Giữ tập sinh ra tất định: đây là ngày của bộ nguồn, không phải
            # thời điểm bất kỳ lúc ai đó chạy lại script.
            "generatedAt": "2026-09-18T00:00:00+07:00",
            "containsTaxpayerIdentity": False,
            "unit": "VND",
        },
        "city2026": {
            "plan": 613_500_000_000_000,
            "augustEstimate": 28_525_000_000_000,
            "ytdAugustEstimate": 489_525_000_000_000,
            "augustActual": 35_492_000_000_000,
            "ytdAugustActual": 496_331_000_000_000,
        },
        "cityHistory": city_history(),
        "forecast2026": forecast_2026(),
        "locationPlans2026": location_rows,
        "locationPlanTotal": sum(int(row["plan"]) for row in location_rows),
        "directory": directory_aggregates(),
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(
        f"Đã ghi {OUTPUT}: {len(location_rows)} địa bàn, "
        f"{payload['directory']['total']:,} bản ghi danh bạ đã tổng hợp."
    )


if __name__ == "__main__":
    main()
