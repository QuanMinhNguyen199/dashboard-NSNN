# -*- coding: utf-8 -*-
"""
Chuyển bộ Excel trong general_data/ thành JSON cho Web quản lý đọc lúc chạy.

Vì sao không nhúng thẳng vào mã nguồn: bộ này có khoảng 21 nghìn dòng người nộp
thuế kèm mã số thuế và tên doanh nghiệp THẬT. Nhúng vào `src/data/` thì dữ liệu
đi theo mọi bản build và lên remote ngay lần push đầu. Ở đây nó được ghi ra
`public/du-lieu-that/` — thư mục đã nằm trong .gitignore — và app nạp bằng
fetch. Không có file thì app rơi về bộ dữ liệu mô phỏng có sẵn.

Chạy: python scripts/nap-du-lieu-that.py
"""
import glob, io, json, os, sys

# Console Windows mac dinh cp1258 va vo khi in chu co dau.
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

try:
    import openpyxl
except ImportError:
    sys.exit("Thiếu openpyxl. Cài bằng: pip install openpyxl")

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NGUON = os.path.join(GOC, "general_data")
DICH = os.path.join(GOC, "public", "du-lieu-that")


def so(v):
    """Ô số: trả float, ô rỗng hoặc lỗi công thức trả None."""
    if v is None or isinstance(v, str):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def chu(v):
    return None if v is None else str(v).strip() or None


def bang(ws, tu_dong, cot):
    """Đọc các dòng có dữ liệu từ `tu_dong`, ánh xạ theo `cot` = {ten: chi_so}."""
    ra = []
    for r in ws.iter_rows(min_row=tu_dong, values_only=True):
        if not any(c is not None for c in r):
            continue
        d = {}
        for ten, i in cot.items():
            if i >= len(r):
                continue
            # "chuong" là MÃ chương ngân sách ("555", "830"), Excel lưu dạng chữ —
            # ép sang số thì `so()` trả None và cả cột biến mất. Mã số thuế và mã
            # CQT cũng vậy: giữ nguyên chữ để không mất số 0 ở đầu.
            d[ten] = so(r[i]) if ten.startswith(("so", "tien", "no", "tong", "ty", "tang", "nguong")) else chu(r[i])
        ra.append(d)
    return ra


def main():
    tep = sorted(glob.glob(os.path.join(NGUON, "*.xlsx")))
    if not tep:
        sys.exit(f"Không thấy file .xlsx nào trong {NGUON}")
    wb = openpyxl.load_workbook(tep[0], read_only=True, data_only=True)
    os.makedirs(DICH, exist_ok=True)

    # ── Quy tắc nguồn: chính là phần khai "cách lấy dữ liệu" của phòng ──
    quy_tac = []
    for r in wb["QuyTac_Nguon"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            quy_tac.append({"noiDung": chu(r[0]), "giaTri": chu(r[1])})
    tra_cuu = {q["noiDung"]: q["giaTri"] for q in quy_tac}

    # ── Nợ theo phòng (ĐVT: TRIỆU ĐỒNG, theo tiêu đề sheet) ──
    ws = wb["So_Sanh_No"]
    tieu_de = chu(list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0][0]) or ""
    moc = ["tongCong", "noKNT", "khoThu", "dangXuLy"]
    don_vi = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[1]:
            continue
        don_vi.append({
            "ten": chu(r[1]),
            "stt": chu(r[0]),
            "laTongHop": r[0] is None,
            "hienTai": {k: so(r[2 + i]) for i, k in enumerate(moc)},
            "soVoiDauNam": {k: so(r[6 + i]) for i, k in enumerate(moc)},
            "soVoiDauNamPhanTram": {k: so(r[10 + i]) for i, k in enumerate(moc)},
            "soVoiThangTruoc": {k: so(r[14 + i]) for i, k in enumerate(moc)},
            "soVoiTuanTruoc": {k: so(r[22 + i]) for i, k in enumerate(moc)},
        })

    # ── Cưỡng chế và tạm hoãn xuất cảnh: cùng khuôn 8 cột ──
    def tong_hop_8(ten_sheet):
        ra = []
        for r in wb[ten_sheet].iter_rows(min_row=4, values_only=True):
            if not r[0]:
                continue
            ra.append({
                "donVi": chu(r[0]),
                "phaiNNT": so(r[1]), "phaiTien": so(r[2]),
                "daNNT": so(r[3]), "daTien": so(r[4]),
                "chuaNNT": so(r[5]), "chuaTien": so(r[6]),
                "tyLe": so(r[7]),
            })
        return ra

    tam_hoan_tong = []
    for r in wb["Bao_cao_tong_hop"].iter_rows(min_row=4, values_only=True):
        if not r[0]:
            continue
        tam_hoan_tong.append({
            "donVi": chu(r[0]), "nntTrangThai06": so(r[1]), "tongNoKhongHoatDong": so(r[2]),
            "daTamHoanNNT": so(r[3]), "daTamHoanTien": so(r[4]), "tyLeDaTamHoan": so(r[5]),
            "chuaTamHoanNNT": so(r[6]), "chuaTamHoanTien": so(r[7]), "tyLeChuaTamHoan": so(r[8]),
        })

    tong_hop = {
        "tenTep": os.path.basename(tep[0]),
        "tieuDe": tieu_de,
        "ngayBaoCao": tra_cuu.get("Ngày báo cáo"),
        "nguongTangNoKNT": so(tra_cuu.get("Ngưỡng tăng nợ KNT")) or tra_cuu.get("Ngưỡng tăng nợ KNT"),
        "donViTien": "triệu đồng cho bảng nợ theo phòng; đồng cho mọi danh sách người nộp thuế",
        "quyTacNguon": quy_tac,
        "noTheoDonVi": don_vi,
        "cuongChe": tong_hop_8("Danh gia Ket qua cuong che"),
        "tamHoanXuatCanh": tong_hop_8("Danh gia Tam hoan XC"),
        "tamHoanTongHop": tam_hoan_tong,
    }

    danh_sach = {
        "no-tang-500.json": bang(wb["DS_DN_TangnoTren500tr"], 3, {
            "stt": 0, "mst": 1, "ten": 2, "noHienTai": 3, "noDauNam": 4, "tangGiam": 5,
            "donVi": 6, "maCQT": 7, "loaiNNT": 8}),
        "chua-cuong-che.json": bang(wb["DS NNT chua cuong che"], 2, {
            "stt": 0, "mst": 1, "ten": 2, "maCQT": 3, "donVi": 4, "loaiNNT": 5, "chuong": 6,
            "noThang": 7, "noNgay": 8, "noDanhGia": 9, "tongNoDanhGia": 10, "nguong": 11,
            "tinhTrang": 12, "bienPhap": 13, "quyetDinh": 14, "ketLuan": 15, "ghiChu": 16}),
        "chua-thxc.json": bang(wb["DS tren 500tr chua Hoan XC"], 3, {
            "stt": 0, "mst": 1, "ten": 2, "maCQT": 3, "donVi": 4, "loaiNNT": 5, "chuong": 6,
            "noThang": 7, "noNgay": 8, "noDanhGia": 9, "tongNoDanhGia": 10, "nguong": 11,
            "tinhTrang": 12, "ngayTamHoan": 13, "ketLuan": 14, "ghiChu": 15}),
        "chua-tam-hoan.json": bang(wb["DS_chua_tam_hoan"], 2, {
            "stt": 0, "mst": 1, "ten": 2, "maCQT": 3, "chuong": 4, "tongNoKhongHoatDong": 5,
            "nhomXuLy": 6, "tinhTrang": 7, "ngayTamHoan": 8, "donVi": 9, "loaiNNT": 10}),
    }
    tong_hop["soDong"] = {k.replace(".json", ""): len(v) for k, v in danh_sach.items()}

    ghi = [("tong-hop.json", tong_hop)] + list(danh_sach.items())
    for ten, noi_dung in ghi:
        duong = os.path.join(DICH, ten)
        with io.open(duong, "w", encoding="utf-8") as f:
            json.dump(noi_dung, f, ensure_ascii=False, separators=(",", ":"))
        print(f"  {ten:24} {os.path.getsize(duong) / 1024:9.1f} kB")
    print(f"Đã ghi {len(ghi)} tệp vào {DICH}")


if __name__ == "__main__":
    main()
